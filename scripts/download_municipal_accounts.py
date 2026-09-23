#!/usr/bin/env python3
"""Download and normalize municipality accounts published by Sambandið."""
from __future__ import annotations

import argparse
import json
import re
import tempfile
from collections import OrderedDict
from pathlib import Path

import openpyxl
import requests

SOURCE_PAGE = "https://www.samband.is/arsreikningar"
CURRENT = {
    "accounts": "https://samband-islenskra-sveitarfelaga.cdn.prismic.io/samband-islenskra-sveitarfelaga/8Z7TVIcrq9Kw-4gI_Net_%C3%81rsreikningar_Pivot.xlsx",
    "per_capita": "https://samband-islenskra-sveitarfelaga.cdn.prismic.io/samband-islenskra-sveitarfelaga/Dxg4F7pQFCDydPRg_Net_Sundurli%C3%B0un_a%C3%B0alsj%C3%B3%C3%B0s_kr_%C3%A1_%C3%ADb%C3%BAa.xlsx",
}
ARCHIVES = {
    2024: "https://samband-islenskra-sveitarfelaga.cdn.prismic.io/samband-islenskra-sveitarfelaga/ae9ggMBOoF08xVBg_%C3%81rb%C3%B3k2025t%C3%B6flur.xlsx",
    2023: "https://samband-islenskra-sveitarfelaga.cdn.prismic.io/samband-islenskra-sveitarfelaga/Z9AQgBsAHJWomUYU_%C3%81rb%C3%B3k2024t%C3%B6flur.xlsx",
    2022: "https://samband-islenskra-sveitarfelaga.cdn.prismic.io/samband-islenskra-sveitarfelaga/Z1GpaZbqstJ98Ehj_%C3%81rb%C3%B3k2023t%C3%B6flur.xlsx",
    2021: "https://samband-islenskra-sveitarfelaga.cdn.prismic.io/samband-islenskra-sveitarfelaga/Z1GpepbqstJ98Ehl_%C3%81rb%C3%B3k2022t%C3%B6flur.xlsx",
    2020: "https://samband-islenskra-sveitarfelaga.cdn.prismic.io/samband-islenskra-sveitarfelaga/Z1Gpj5bqstJ98Ehn_%C3%81rb%C3%B3k2021t%C3%B6flur.xlsx",
    2019: "https://samband-islenskra-sveitarfelaga.cdn.prismic.io/samband-islenskra-sveitarfelaga/Z1GwapbqstJ98Em9_%C3%81rb%C3%B3k2020t%C3%B6flur.xlsx",
}


def clean_name(value) -> str:
    return re.sub(r"^\d+\s+", "", str(value or "")).strip()


def is_municipality(name: str) -> bool:
    return bool(name) and name.lower() not in {"grand total", "landið allt"}


def metric_key(prefix: str, label: str) -> str:
    safe = re.sub(r"[^a-z0-9]+", "-", label.lower().replace("ð", "d").replace("þ", "th"))
    return f"{prefix}-{safe.strip('-')}"


def add_metric(store, prefix, label, section, year, values, per_capita=False):
    key = metric_key(prefix, label)
    item = store.setdefault(key, {"key": key, "label": label, "section": section,
                                  "per_capita": per_capita, "percentage": False,
                                  "values": {}})
    item["values"][str(year)] = values


def parse_current(accounts_path: Path, pc_path: Path, store, municipality_names):
    ws = openpyxl.load_workbook(accounts_path, data_only=True).active
    year = int(ws.cell(2, 2).value)
    columns = [(col, clean_name(ws.cell(6, col).value)) for col in range(4, ws.max_column + 1)
               if ws.cell(6, col).value and is_municipality(clean_name(ws.cell(6, col).value))]
    section = None
    for row in range(7, ws.max_row + 1):
        if ws.cell(row, 1).value:
            section = str(ws.cell(row, 1).value).strip()
        label = ws.cell(row, 2).value
        if not label:
            continue
        values = {name: float(ws.cell(row, col).value) for col, name in columns
                  if isinstance(ws.cell(row, col).value, (int, float))}
        if values:
            municipality_names.update(values)
            add_metric(store, "acc", str(label).strip(), section, year, values)

    ws = openpyxl.load_workbook(pc_path, data_only=True).active
    pc_year = int(ws.cell(2, 2).value)
    columns = [(col, clean_name(ws.cell(5, col).value)) for col in range(4, ws.max_column + 1)
               if ws.cell(5, col).value and is_municipality(clean_name(ws.cell(5, col).value))]
    for row in range(6, ws.max_row + 1):
        label = ws.cell(row, 1).value
        if not label:
            continue
        values = {name: float(ws.cell(row, col).value) for col, name in columns
                  if isinstance(ws.cell(row, col).value, (int, float))}
        if values:
            municipality_names.update(values)
            add_metric(store, "pc", str(label).strip(), "Málaflokkar", pc_year, values, True)
    return year


def parse_archive(path: Path, expected_year: int, store, municipality_names):
    book = openpyxl.load_workbook(path, data_only=True)
    ws = book["Tafla 6"]
    year = int(ws.cell(3, 1).value or expected_year)
    last_name = ""
    columns = []
    for col in range(2, ws.max_column + 1):
        if ws.cell(5, col).value:
            last_name = clean_name(ws.cell(5, col).value)
        if col >= 4 and ws.cell(8, col).value == "A hluti" and last_name:
            columns.append((col, last_name))
    section = None
    for row in range(9, ws.max_row + 1):
        label = ws.cell(row, 1).value
        if not label:
            continue
        values = {name: float(ws.cell(row, col).value) for col, name in columns
                  if isinstance(ws.cell(row, col).value, (int, float))}
        if values:
            municipality_names.update(values)
            add_metric(store, "acc", str(label).strip(), section or "Ársreikningur", year, values)
        else:
            section = str(label).split("(")[0].strip()

    ws = book["Tafla 8"]
    for row in range(1, ws.max_row + 1):
        category, name = ws.cell(row, 1).value, ws.cell(row, 4).value
        net_per_capita = ws.cell(row, 15).value
        if category and name and isinstance(net_per_capita, (int, float)):
            label = str(category).strip()
            municipality = clean_name(name)
            key = metric_key("pc", label)
            if key not in store:
                add_metric(store, "pc", label, "Málaflokkar", year, {}, True)
            store[key]["values"].setdefault(str(year), {})[municipality] = -float(net_per_capita)


def download(url: str, path: Path):
    response = requests.get(url, timeout=90)
    response.raise_for_status()
    path.write_bytes(response.content)


def add_derived_metrics(store):
    """Add size-neutral indicators calculated from the published account lines."""
    revenue = store.get("acc-tekjur")
    result = store.get("acc-rekstrarnidurstada")
    if not revenue or not result:
        return
    values = {}
    for year in sorted(set(revenue["values"]) & set(result["values"])):
        ratios = {}
        for name, amount in result["values"][year].items():
            denominator = revenue["values"][year].get(name)
            if denominator:
                ratios[name] = amount / denominator * 100
        if ratios:
            values[year] = ratios
    store["ratio-operating-result"] = {
        "key": "ratio-operating-result",
        "label": "Rekstrarniðurstaða sem hlutfall af tekjum",
        "section": "Lykiltölur",
        "per_capita": False,
        "percentage": True,
        "values": values,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path("static/municipal_accounts.json"))
    args = parser.parse_args()
    store, municipalities = OrderedDict(), set()
    with tempfile.TemporaryDirectory(prefix="municipal-accounts-") as tmp:
        folder = Path(tmp)
        current_accounts, current_pc = folder / "current.xlsx", folder / "current-pc.xlsx"
        download(CURRENT["accounts"], current_accounts)
        download(CURRENT["per_capita"], current_pc)
        current_year = parse_current(current_accounts, current_pc, store, municipalities)
        for year, url in ARCHIVES.items():
            path = folder / f"archive-{year}.xlsx"
            download(url, path)
            parse_archive(path, year, store, municipalities)
    add_derived_metrics(store)
    payload = {"source": SOURCE_PAGE,
               "years": sorted({int(y) for m in store.values() for y in m["values"]}, reverse=True),
               "current_year": current_year,
               "municipalities": sorted(municipalities), "metrics": list(store.values())}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Wrote {args.output}: {len(payload['municipalities'])} municipalities, "
          f"{len(payload['metrics'])} metrics, years {min(payload['years'])}–{max(payload['years'])}")


if __name__ == "__main__":
    main()
